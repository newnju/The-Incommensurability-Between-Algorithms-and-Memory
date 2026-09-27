# -*- coding: utf-8 -*-
"""Data loader: reads ALL figure data directly from the coding workbook.

The workbook coding_data/per_image_coding_table_v1.0.xlsx is the single
source of truth (ground truth). No figure values are hard-coded.

Sheets used:
- 稳定性_留一: Bootstrap selection frequency pi_j and leave-one-model-out
  inclusion matrix (f1..f18 x M1..M6).
- 维度_Kappa: dimension-level linear-weighted kappa with 95% CI, pooled
  kappa, and overall criterion-Y kappa.
- 提示条件汇总: dimension means per condition and grade-2 proportions.
- 编码员A / 编码员B: per-image atomic indicators (condition contrasts for
  Fig. 4b are recomputed from these with the preregistered bootstrap:
  B=CONTRAST_B, seed 20260824, stratified by model, consensus = mean of
  coders, NA -> 0; see CONTRAST_B below for why the contrast CIs use a
  larger B than the selection-frequency bootstrap).
"""
import os
import warnings

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBOOK = os.path.join(HERE, '..', 'coding_data',
                        'per_image_coding_table_v1.0.xlsx')

BOOT_B = 10000          # bootstrap replicates for the selection frequency pi
BOOT_SEED = 20260824    # preregistered seed, shared by every random process

# Bootstrap replicates for the Fig. 4b condition contrasts.
#
# A contrast is a difference between two means of 54 consensus scores, and
# every consensus score lies on the 0.5 grid (it is the mean of two integer
# ratings).  The bootstrap distribution of the contrast is therefore
# *discrete*, with spacing 0.5/54 = 0.00926.  The 97.5th percentile of such
# a distribution does not converge smoothly at small B: it hops between
# adjacent mass points from one RNG stream to the next.  Measured over 200
# seeds at B=10000, the D1 P2-P1 upper limit alternates between +0.1944
# (125/200) and +0.2037 (71/200); the D5 P2-P1 lower limit alternates
# between -0.0833 (122/200) and -0.0741 (72/200); the D1 P3-P2 upper limit
# between +0.3889 (108/200) and +0.3796 (90/200); the D5 P3-P2 lower limit
# between +0.0370 (134/200) and +0.0278 (60/200).  At that B the third
# decimal is a property of the seed, not of the data - and for D1 P2-P1 it
# decides whether the upper limit "slightly exceeds" +-0.20 or not.
#
# At B=1000000 every endpoint is bit-identical across streams, so the
# reported limits are reproducible.  The seed is still the preregistered
# 20260824; at this B it no longer affects the answer.
CONTRAST_B = 1000000


def _load_workbook():
    warnings.filterwarnings('ignore')
    return openpyxl.load_workbook(WORKBOOK, data_only=True)


def stability_data():
    """Returns (labels, pi, matrix) from sheet 稳定性_留一."""
    wb = _load_workbook()
    ws = wb['稳定性_留一']
    labels, pi, mat = [], [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None or not str(r[0]).startswith('f'):
            break
        labels.append(str(r[0]))
        pi.append(float(r[2]))
        mat.append([int(x) for x in r[3:9]])
    wb.close()
    return labels, np.array(pi), np.array(mat, dtype=int)


def kappa_data():
    """Returns dict with per-dimension kappa, pooled kappa, criterion-Y."""
    wb = _load_workbook()
    ws = wb['维度_Kappa']
    dims = {}
    pooled = crity = None
    for r in ws.iter_rows(min_row=1, values_only=True):
        key = r[0]
        if key is None:
            continue
        key = str(key)
        if key.startswith(('D1', 'D2', 'D3', 'D4', 'D5')):
            short = key.split()[0]          # 'D1' ... 'D5'
            dims[short] = (float(r[5]), float(r[6]), float(r[7]))
        elif key.startswith('微观合并'):
            pooled = (float(r[5]), float(r[6]), float(r[7]))
        elif key == '线性加权Kappa':        # criterion-Y row
            crity = (float(r[1]), float(r[2]), float(r[3]))
    wb.close()
    return {'dims': dims, 'pooled': pooled, 'criterion_Y': crity}


def condition_summary():
    """Returns (conditions, dim_means, grade2) from sheet 提示条件汇总.

    dim_means: {dim_short: [P1, P2, P3]} for D1..D5
    grade2: list of (condition, k, n) grade-2 counts.
    """
    wb = _load_workbook()
    ws = wb['提示条件汇总']
    rows = list(ws.iter_rows(values_only=True))
    conds, g2 = [], []
    for r in rows:
        if r[0] in ('P1', 'P2', 'P3') and isinstance(r[2], (int, float)):
            # Y=0..Y=2 counts at idx2..6, effective N at idx7
            g2.append((r[0], int(r[6]), int(r[7])))
    dim_means = {}
    for r in rows:
        if r[0] is None:
            continue
        key = str(r[0])
        if key[:2] in ('D1', 'D2', 'D3', 'D4', 'D5') and \
                all(isinstance(v, (int, float)) for v in r[1:4]):
            dim_means[key[:2]] = [float(r[1]), float(r[2]), float(r[3])]
    conds = ['P1', 'P2', 'P3']
    wb.close()
    return conds, dim_means, g2


def _consensus_recs():
    """Per-image records: (model, condition, D1, D5) with consensus scoring
    (mean of coder A and B, NA -> 0) for dimensions D1=f2 and D5=f18."""
    wb = _load_workbook()
    recs = {}
    for sheet in ('编码员A', '编码员B'):
        ws = wb[sheet]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            v = lambda x: 0.0 if x is None or str(x) == 'NA' else float(x)
            rec = recs.setdefault(str(r[0]), {'model': r[1], 'cond': r[3]})
            rec.setdefault('vals', []).append((v(r[7]), v(r[23])))
    wb.close()
    out = []
    for iid in sorted(recs):
        rec = recs[iid]
        a, b = rec['vals']
        out.append((rec['model'], rec['cond'], (a[0] + b[0]) / 2,
                    (a[1] + b[1]) / 2))
    return out


def condition_contrasts():
    """Fig. 4b contrasts: {label: (point, lo, hi)} via the preregistered
    model-stratified bootstrap (B=CONTRAST_B, seed 20260824).

    Design: 18 model x condition cells are resampled with replacement
    within cell, the resampled consensus scores are pooled over the six
    models, and the contrast is the difference of the two pooled means.
    Percentile interval.  See CONTRAST_B for why B is large here.
    """
    recs = _consensus_recs()
    models = sorted({r[0] for r in recs})
    vals = np.array([r[2:4] for r in recs], dtype=float)
    conds = np.array([r[1] for r in recs])
    model_of = np.array([r[0] for r in recs])
    idx = {m: {c: np.where((model_of == m) & (conds == c))[0]
               for c in ('P1', 'P2', 'P3')}
           for m in models}
    n_per_cond = sum(idx[m]['P1'].size for m in models)
    rng = np.random.default_rng(BOOT_SEED)

    def boot(c1, c2, dim):
        acc = {c1: np.zeros(CONTRAST_B), c2: np.zeros(CONTRAST_B)}
        for m in models:
            for c in (c1, c2):
                col = vals[idx[m][c], dim]
                draw = rng.choice(col.size, size=(CONTRAST_B, col.size),
                                  replace=True)
                acc[c] += col[draw].sum(axis=1)
        diffs = (acc[c2] - acc[c1]) / n_per_cond
        point = vals[conds == c2, dim].mean() - \
            vals[conds == c1, dim].mean()
        lo, hi = np.percentile(diffs, [2.5, 97.5])
        return point, lo, hi

    out = {}
    for dim, dname in ((0, 'D1'), (1, 'D5')):
        for c1, c2 in (('P1', 'P2'), ('P2', 'P3')):
            out[f'{dname} {c2}-{c1}'] = boot(c1, c2, dim)
    return out


def criterion_y_by_id():
    """Returns {ImageID: consensus criterion Y} from sheet 共识效标Y.

    Used by Fig. 3, whose panels annotate the criterion Y of the displayed
    image; reading it here keeps that annotation tied to the workbook
    instead of hard-coded in the plotting script.
    """
    wb = _load_workbook()
    out = {}
    for r in wb['共识效标Y'].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        out[str(r[0])] = float(r[7])          # column H = consensus Y
    wb.close()
    return out


def full_dataset():
    """Full modelling dataset used by the stability analysis.

    Returns (X, y, meta):
    - X: (162, 18) feature matrix — mean of coder A and B per atomic indicator
      (NA -> 0), columns = f1..f18
    - y: (162,) consensus criterion Y (sheet 共识效标Y, column H)
    - meta: list of (model, condition) per image

    Feature construction (coder mean, NA -> 0) and the consensus-Y rule follow
    the preregistered specification; nothing is hard-coded here.
    """
    wb = _load_workbook()
    feats = {}
    for sheet in ('编码员A', '编码员B'):
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            v = lambda x: 0.0 if x is None or str(x) == 'NA' else float(x)
            rec = feats.setdefault(str(r[0]),
                                  {'model': r[1], 'cond': r[3], 'v': []})
            rec['v'].append([v(x) for x in r[6:24]])      # f1..f18
    cons = {}
    for r in wb['共识效标Y'].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        cons[str(r[0])] = float(r[7])                     # consensus Y
    wb.close()

    X, y, meta = [], [], []
    for iid in sorted(feats):
        a, b = feats[iid]['v']
        X.append([(a[j] + b[j]) / 2 for j in range(18)])
        y.append(cons[iid])
        meta.append((feats[iid]['model'], feats[iid]['cond']))
    return np.array(X), np.array(y), meta
